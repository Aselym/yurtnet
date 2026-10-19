"""Excel / CSV veri kaynagi.

Kullanicinin yukledigi tabloyu, onayladigi kolon eslemesine gore Logo
kaynagiyla AYNI arayuzle verir (urunler, cariler, temsilciler, satirlar...).
Boylece ambar, FIFO, lisans takibi ve tum analizler degismeden calisir.
"""

import csv
import hashlib
import io
import os
import zlib

import esleme

# Yon -> Logo hareket kodlari (TRCODE, IOCODE). warehouse.yon_bul bunlari geri cevirir.
YON_KODLARI = {'SATIS': (8, 4), 'SATIS_IADE': (3, 1), 'ALIS': (1, 1), 'ALIS_IADE': (6, 4)}


def kararli_ref(sayac, fis, urun_anahtari):
    """Satir kimligi dosyadaki siraya degil icerige bagli: fatura no + urun + faturadaki tekrar sirasi.
    Ayni satir her yuklemede (ya da Excel <-> e-Fatura arasinda) ayni kimligi alir; lisanslar cogalmaz."""
    anahtar = f'{fis}|{urun_anahtari}'
    sayac[anahtar] = sayac.get(anahtar, 0) + 1
    # 48 bit: on binlerce satirda bile carpisma olasiligi ihmal edilebilir
    return int(hashlib.sha1(f'{anahtar}|{sayac[anahtar]}'.encode('utf-8')).hexdigest()[:12], 16) + 1


def tablo_oku(yol, sayfa=None, en_fazla=None):
    """-> (sayfa_adlari, secilen_sayfa, basliklar, satirlar). Baslik satirini kendisi bulur."""
    uzanti = os.path.splitext(yol)[1].lower()
    if uzanti in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(yol, read_only=True, data_only=True)
        adlar = wb.sheetnames
        ws = wb[sayfa] if sayfa in adlar else wb[adlar[0]]
        ham = []
        for r in ws.iter_rows(values_only=True):
            ham.append(list(r))
            if en_fazla and len(ham) >= en_fazla + 20:
                break
        secilen = ws.title
        wb.close()
    elif uzanti in ('.csv', '.txt', '.tsv'):
        with open(yol, 'rb') as fh:
            veri = fh.read()
        metin = None
        for kod in ('utf-8-sig', 'cp1254', 'latin-1'):
            try:
                metin = veri.decode(kod)
                break
            except UnicodeDecodeError:
                continue
        try:
            ayirac = csv.Sniffer().sniff(metin[:5000], delimiters=';,\t|').delimiter
        except csv.Error:
            ayirac = ';' if metin[:2000].count(';') > metin[:2000].count(',') else ','
        ham = [r for r in csv.reader(io.StringIO(metin), delimiter=ayirac)]
        if en_fazla:
            ham = ham[:en_fazla + 20]
        adlar, secilen = ['CSV'], 'CSV'
    else:
        raise ValueError('Desteklenen dosya türleri: .xlsx, .csv')

    # Baslik satiri: ilk 20 satirda en cok dolu metin hucresi olan satir
    # (raporlarin ustunde sirket adi, tarih gibi satirlar olabiliyor)
    en_iyi, puan = 0, -1
    for i, r in enumerate(ham[:20]):
        p = sum(1 for v in r if isinstance(v, str) and v.strip() and esleme.sayi(v) is None)
        if p > puan:
            en_iyi, puan = i, p
    basliklar = [str(v).strip() if v not in (None, '') else f'Kolon {j + 1}' for j, v in enumerate(ham[en_iyi] if ham else [])]
    satirlar = [r for r in ham[en_iyi + 1:] if any(v not in (None, '') for v in r)]
    if en_fazla:
        satirlar = satirlar[:en_fazla]
    return adlar, secilen, basliklar, satirlar


class DosyaKaynagi:
    """Logo kaynagi taklidi. esleme_ayari: {'dosya', 'sayfa', 'esleme': {alan: kolon_index}, 'varsayilan_yon'}."""

    mod = 'dosya'

    def __init__(self, esleme_ayari, klasor):
        self.ayar = esleme_ayari
        self.yol = os.path.join(klasor, esleme_ayari['dosya'])
        self._cozuldu = False
        self._markalar = {}

    def __enter__(self):
        self._coz()
        return self

    def __exit__(self, *a):
        return False

    # -- kaynak arayuzu
    def markalar(self):
        return dict(self._markalar)

    def urunler(self):
        return self._urunler

    def hizmetler(self):
        return []

    def cariler(self):
        return self._cariler

    def temsilciler(self):
        return self._temsilciler

    def doviz_kurlari(self):
        return {}

    def satirlar(self, baslangic=None):
        for s in self._satirlar:
            if not baslangic or s['date'] >= baslangic:
                yield s

    def fatura_sayisi(self):
        return len({s['fiche_no'] for s in self._satirlar})

    # -- donusum
    def _coz(self):
        if self._cozuldu:
            return
        _, _, basliklar, satirlar = tablo_oku(self.yol, self.ayar.get('sayfa'))
        e = {k: int(v) for k, v in (self.ayar.get('esleme') or {}).items() if v is not None and v != ''}
        varsayilan = self.ayar.get('varsayilan_yon') or 'SATIS'

        def al(r, alan):
            i = e.get(alan)
            return r[i] if i is not None and i < len(r) else None

        urunler, cariler, temsilciler = {}, {}, {}
        markalar = {}
        cikti, atlanan, refs = [], 0, {}
        for n, r in enumerate(satirlar, 1):
            gun = esleme.tarih(al(r, 'tarih'))
            urun_adi = str(al(r, 'urun') or '').strip()
            if not gun or not urun_adi:
                atlanan += 1
                continue
            miktar = esleme.sayi(al(r, 'miktar'))
            bf = esleme.sayi(al(r, 'birim_fiyat'))
            tutar = esleme.sayi(al(r, 'tutar'))
            if tutar is None and bf is not None:
                tutar = bf * (miktar if miktar is not None else 1)
            if tutar is None:
                atlanan += 1
                continue
            if miktar is None or miktar == 0:
                miktar = 1.0
            yon = esleme.yon_coz(al(r, 'yon'), tutar, varsayilan)
            tutar, miktar = abs(tutar), abs(miktar)
            usd = esleme.sayi(al(r, 'tutar_usd'))
            kur = esleme.sayi(al(r, 'kur'))
            if (not kur or kur <= 0) and usd:
                kur = tutar / abs(usd) if usd else None

            ukod = str(al(r, 'urun_kodu') or '').strip()
            uanahtar = ukod or urun_adi.upper()
            if uanahtar not in urunler:
                urunler[uanahtar] = {'ref': len(urunler) + 1, 'code': ukod, 'name': urun_adi,
                                     'grp': str(al(r, 'kategori') or '').strip(), 'markref': None}
                marka = str(al(r, 'marka') or '').strip()
                if marka:
                    markalar[len(urunler)] = marka
            cari_adi = str(al(r, 'cari') or '').strip() or 'Bilinmeyen cari'
            ckod = str(al(r, 'cari_kodu') or '').strip()
            canahtar = ckod or cari_adi.upper()
            if canahtar not in cariler:
                cariler[canahtar] = {'ref': len(cariler) + 1, 'code': ckod, 'name': cari_adi, 'cardtype': None,
                                     'email': '', 'phone': '', 'city': str(al(r, 'sehir') or '').strip()}
            tem = str(al(r, 'temsilci') or '').strip()
            if tem and tem not in temsilciler:
                temsilciler[tem] = {'ref': len(temsilciler) + 1, 'code': '', 'name': tem}
            trcode, iocode = YON_KODLARI[yon]
            fis = str(al(r, 'fatura_no') or '').strip() or f'{gun}-{canahtar[:20]}'
            cikti.append({
                'period': 'K', 'line_ref': kararli_ref(refs, fis, uanahtar), 'inv_ref': zlib.crc32(fis.encode('utf-8')) % 2_000_000_000 or 1, 'date': gun,
                'trcode': trcode, 'iocode': iocode, 'linetype': 0,
                'item_ref': urunler[uanahtar]['ref'], 'client_ref': cariler[canahtar]['ref'],
                'salesman_ref': temsilciler[tem]['ref'] if tem else 0,
                'qty': miktar, 'price': tutar / miktar if miktar else tutar, 'gross': tutar, 'net': tutar,
                'rate': kur, 'line_cancelled': 0, 'fiche_no': fis, 'inv_cancelled': 0,
                'inv_salesman': 0, 'inv_client': 0,
            })
        # marka bilgisi dosyada varsa Logo'daki marka karti gibi davranir
        self._markalar = markalar
        for u in urunler.values():
            u['markref'] = u['ref'] if u['ref'] in markalar else None
        self._urunler = list(urunler.values())
        self._cariler = list(cariler.values())
        self._temsilciler = list(temsilciler.values())
        self._satirlar = cikti
        self.atlanan = atlanan
        self._cozuldu = True
