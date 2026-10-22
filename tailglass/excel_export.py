"""Excel'e aktarma: ekranda ne goruluyorsa (ayni filtre ve kirilimlarla) o.

Gorunum: Tailglass marka renkleri, baslik bandi, seritli satirlar, belirgin toplam,
izgara cizgisi yok, yazdirmaya hazir (yatay, sayfa genisligine sigar).
Anlam: donemler Turkce ("Haziran 2026"), pay ve onceki donemle karsilastirma
sutunlari, Excel'in kendi grafigi, ilk sayfada ozet. Kurus/sent gosterilmez.
"""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

import analytics

# ------------------------------------------------------------------ tema
MUREKKEP, TURKUAZ, TURKUAZ_KOYU = '0B0F0E', '1FC8B6', '0F8F82'
GRI, ACIK_GRI, SERIT, TOPLAM_ZEMIN = '5B6866', 'D9E1DF', 'F5F8F7', 'E6F7F5'
YAZI = 'Calibri'

BICIM = {'tl': '#,##0 "₺";-#,##0 "₺"', 'usd': '"$"#,##0;-"$"#,##0', 'num': '#,##0', 'pct': '0.0"%"',
         'degisim': '+0.0"%";-0.0"%";0.0"%"', 'puan': '+0.0" p";-0.0" p";0.0" p"', 'date': '@', 'text': '@'}

F_MARKA = Font(name=YAZI, size=9, color=TURKUAZ_KOYU, bold=True)
F_BASLIK = Font(name=YAZI, size=16, bold=True, color=MUREKKEP)
F_ALT = Font(name=YAZI, size=10, color=GRI)
F_KOLON = Font(name=YAZI, size=10, bold=True, color='FFFFFF')
F_HUCRE = Font(name=YAZI, size=10, color=MUREKKEP)
F_TOPLAM = Font(name=YAZI, size=10, bold=True, color=MUREKKEP)
F_ETIKET = Font(name=YAZI, size=10, bold=True, color=MUREKKEP)
F_POZ = Font(name=YAZI, size=10, color='15803D')
F_NEG = Font(name=YAZI, size=10, color='B91C1C')
D_KOLON = PatternFill('solid', fgColor=MUREKKEP)
D_SERIT = PatternFill('solid', fgColor=SERIT)
D_TOPLAM = PatternFill('solid', fgColor=TOPLAM_ZEMIN)
INCE = Side(style='thin', color=ACIK_GRI)
KALIN = Side(style='medium', color=MUREKKEP)
TURKUAZ_CIZGI = Side(style='medium', color=TURKUAZ)

AYLAR = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım', 'Aralık']
ZAMAN = {'day', 'week', 'month', 'quarter', 'year', 'weekday'}


def _bicim(tip, v):
    if tip == 'num' and isinstance(v, float) and not v.is_integer() and abs(v) < 1000:
        return '#,##0.00'
    return BICIM.get(tip, 'General')


def donem_yazisi(dim, v):
    """Zaman anahtarini okunur Turkceye cevirir: 2026-06 -> Haziran 2026."""
    s = '' if v is None else str(v)
    try:
        if dim == 'month' and len(s) == 7:
            return f"{AYLAR[int(s[5:7]) - 1]} {s[:4]}"
        if dim == 'quarter' and '-Ç' in s:
            y, q = s.split('-Ç')
            return f'{y} {q}. çeyrek'
        if dim == 'week' and '-H' in s:
            y, w = s.split('-H')
            return f'{y} / {int(w)}. hafta'
        if dim == 'day' and len(s) == 10:
            return date.fromisoformat(s).strftime('%d.%m.%Y')
        if dim == 'weekday':
            return s[2:] if s[:1].isdigit() else s
    except (ValueError, IndexError):
        pass
    return s


def _tarih(t):
    try:
        return date.fromisoformat(str(t)[:10]).strftime('%d.%m.%Y')
    except (TypeError, ValueError):
        return t or ''


def _aralik_yazisi(istek):
    return f"{_tarih(istek.get('date_from')) or 'Başlangıç'} – {_tarih(istek.get('date_to')) or 'bugün'}"


def _filtre_yazisi(istek):
    parca = []
    for boyut, degerler in (istek.get('filters') or {}).items():
        onek = ''
        if isinstance(degerler, dict):
            onek = 'hariç ' if degerler.get('not') else ''
            degerler = degerler.get('values') or []
        if degerler:
            ad = analytics.BOYUTLAR.get(boyut, (None, boyut))[1]
            liste = ', '.join(map(str, degerler[:6])) + (f' (+{len(degerler) - 6})' if len(degerler) > 6 else '')
            parca.append(f'{ad}: {onek}{liste}')
    return ' · '.join(parca)


def _sayfa_hazirla(ws, yatay=True):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = 'landscape' if yatay else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = 'Tailglass Analytics'
    ws.oddFooter.right.text = 'Sayfa &P / &N'


def _ust_bant(ws, baslik, alt=(), genislik=8):
    """Marka satiri + baslik + alt bilgiler. Doner: bir sonraki bos satir."""
    ws.cell(row=1, column=1, value='TAILGLASS ANALYTICS').font = F_MARKA
    ws.cell(row=2, column=1, value=baslik).font = F_BASLIK
    ws.row_dimensions[2].height = 24
    r = 3
    for satir in alt:
        if satir:
            ws.cell(row=r, column=1, value=satir).font = F_ALT
            r += 1
    for j in range(1, genislik + 1):
        ws.cell(row=r, column=j).border = Border(bottom=TURKUAZ_CIZGI)
    ws.row_dimensions[r].height = 6
    return r + 2


def _sayfa_yaz(ws, kolonlar, satirlar, toplam=None, baslik=None, alt=(), r0=None, grafik=None):
    """Bicimli tablo. kolonlar: [{'key','label','type'}]. Doner: (baslik_satiri, son_satir)."""
    _sayfa_hazirla(ws)
    if r0 is None:
        r0 = _ust_bant(ws, baslik, alt, max(len(kolonlar), 4)) if baslik else 1
    for j, k in enumerate(kolonlar, 1):
        c = ws.cell(row=r0, column=j, value=k['label'])
        c.fill, c.font = D_KOLON, F_KOLON
        sayisal = k.get('type') in ('tl', 'usd', 'num', 'pct', 'degisim', 'puan')
        c.alignment = Alignment(horizontal='right' if sayisal else 'left', vertical='center', wrap_text=True)
    ws.row_dimensions[r0].height = 30
    for i, satir in enumerate(satirlar, r0 + 1):
        serit = (i - r0) % 2 == 0
        for j, k in enumerate(kolonlar, 1):
            v = satir.get(k['key'])
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_HUCRE
            if serit:
                c.fill = D_SERIT
            c.border = Border(bottom=INCE)
            if isinstance(v, (int, float)) and k.get('type') in BICIM:
                c.number_format = _bicim(k['type'], v)
                if k.get('type') in ('degisim', 'puan'):
                    c.font = F_POZ if v > 0 else F_NEG if v < 0 else F_HUCRE
    son = r0 + len(satirlar)
    if toplam:
        son += 1
        for j, k in enumerate(kolonlar, 1):
            c = ws.cell(row=son, column=j)
            c.fill, c.font, c.border = D_TOPLAM, F_TOPLAM, Border(top=KALIN, bottom=KALIN)
            if j == 1:
                c.value = 'Toplam'
                continue
            v = toplam.get(k['key'])
            if isinstance(v, (int, float)):
                c.value = v
                c.number_format = _bicim(k.get('type'), v)
    for j, k in enumerate(kolonlar, 1):
        olcu = [len(str(k['label'])) * 0.9] + [len(_gorunur(s.get(k['key']), k.get('type'))) for s in satirlar[:400]]
        ws.column_dimensions[get_column_letter(j)].width = min(max(11, max(olcu) + 3), 55 if j == 1 else 24)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2 if len(kolonlar) > 3 else 1)
    if satirlar:
        ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(kolonlar))}{r0 + len(satirlar)}'
    ws.print_title_rows = f'{r0}:{r0}'
    if grafik and len(satirlar) >= 2:
        _grafik_ekle(ws, r0, satirlar, kolonlar, **grafik)
    return r0, son


def _gorunur(v, tip):
    if isinstance(v, float):
        return f'{v:,.0f} ₺' if tip == 'tl' else f'{v:,.1f}'
    return str(v if v is not None else '')


def _grafik_ekle(ws, r0, satirlar, kolonlar, olcu_key, zaman=False, en_fazla=15):
    """Ana olcu icin Excel grafigi (tablonun sagina). Kategoride ilk 15, zamanda tum donemler."""
    j = next((i for i, k in enumerate(kolonlar, 1) if k['key'] == olcu_key), None)
    if not j:
        return
    n = len(satirlar) if zaman else min(len(satirlar), en_fazla)
    g = LineChart() if zaman and n > 18 else BarChart()
    if isinstance(g, BarChart):
        g.type = 'col'
        g.gapWidth = 60
    g.title = kolonlar[j - 1]['label']
    g.style = 2
    g.legend = None
    g.height, g.width = 9, 18
    g.y_axis.numFmt = '#,##0'
    g.y_axis.delete = False
    g.x_axis.delete = False
    g.x_axis.tickLblPos = 'low'
    veri = Reference(ws, min_col=j, min_row=r0, max_row=r0 + n)
    adlar = Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + n)
    g.add_data(veri, titles_from_data=True)
    g.set_categories(adlar)
    s = g.series[0]
    s.graphicalProperties.solidFill = TURKUAZ_KOYU
    s.graphicalProperties.line.solidFill = TURKUAZ_KOYU
    ws.add_chart(g, f'{get_column_letter(len(kolonlar) + 2)}{r0}')


def _bilgi_sayfasi(wb, istek, baslik):
    """Diger disa aktarimlar icin ilk sayfa: rapor bilgisi."""
    ws = wb.active
    ws.title = 'Rapor bilgisi'
    _sayfa_hazirla(ws, yatay=False)
    r = _ust_bant(ws, baslik, (f"Üretildi: {datetime.now().strftime('%d.%m.%Y %H:%M')}",), 2)
    satirlar = [('Dönem', _aralik_yazisi(istek))]
    if istek.get('direction'):
        yonler = {d['key']: d['label'] for d in analytics.meta()['directions']}
        satirlar.append(('Hareket', yonler.get(istek['direction'], istek['direction'])))
    if _filtre_yazisi(istek):
        satirlar.append(('Filtreler', _filtre_yazisi(istek)))
    if istek.get('search'):
        satirlar.append(('Arama', istek['search']))
    if istek.get('group_by'):
        gb = istek['group_by'] if isinstance(istek['group_by'], list) else [istek['group_by']]
        satirlar.append(('Kırılım', ' › '.join(analytics.BOYUTLAR.get(g, (None, g))[1] for g in gb)))
    satirlar.append(('Not', 'Maliyet ve kâr FIFO yöntemiyle tahmin edilmiştir; USD tutarlar işlem günündeki kurla hesaplanmıştır.'))
    for k, v in satirlar:
        ws.cell(row=r, column=1, value=k).font = F_ETIKET
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.alignment = F_HUCRE, Alignment(wrap_text=True, vertical='top')
        r += 1
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 90


def _kaydet(wb):
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def _sayfa_adi(wb, ad):
    ad = ''.join(' ' if ch in '[]:*?/' + chr(92) else ch for ch in ad).strip()[:31] or 'Sayfa'
    aday, n = ad, 2
    while aday in wb.sheetnames:
        aday = f'{ad[:27]} ({n})'
        n += 1
    return aday


# ------------------------------------------------------------------ analiz (kart / odak) tablosu
def analiz_tablosu(istek):
    """Kirilimli sorgu + okunur donemler + pay + onceki esit donemle karsilastirma.
    Doner: {'kolonlar', 'satirlar', 'toplam', 'ozet', 'zaman', 'ana'}"""
    istek = {**istek, 'limit': 0}
    sonuc = analytics.sorgu(istek)
    gb = istek.get('group_by') or []
    olculer = istek.get('measures') or ['net_tl']
    ana = f'm_{olculer[0]}'
    ana_tip = analytics.OLCULER.get(olculer[0], (None, None, 'num'))[2]
    zaman = bool(gb) and gb[0] in ZAMAN
    satirlar = [dict(r) for r in sonuc['rows']]
    if zaman:
        satirlar.sort(key=lambda r: str(r.get('d0')))

    onceki_toplam, once = None, {}
    if istek.get('date_from') and istek.get('date_to'):
        of, ot = analytics.onceki_donem(istek['date_from'], istek['date_to'])
        o = analytics.sorgu({**istek, 'date_from': of, 'date_to': ot})
        onceki_toplam = o['totals'].get(ana)
        once = {(r.get('d0'), r.get('d1')): r.get(ana) for r in o['rows']}

    toplam_ana = abs(sonuc['totals'].get(ana) or 0)
    for i, r in enumerate(satirlar):
        v = r.get(ana)
        if ana_tip != 'pct' and toplam_ana and isinstance(v, (int, float)):
            r['_pay'] = 100.0 * v / toplam_ana
        ov = (satirlar[i - 1].get(ana) if i else None) if (zaman and len(gb) == 1) else once.get((r.get('d0'), r.get('d1')))
        if isinstance(v, (int, float)) and isinstance(ov, (int, float)) and ov:
            r['_onceki'] = ov
            r['_degisim'] = (v - ov) if ana_tip == 'pct' else 100.0 * (v - ov) / abs(ov)
        for n, g in enumerate(gb[:2]):
            r[f'd{n}'] = donem_yazisi(g, r.get(f'd{n}'))

    kolonlar = [c for c in sonuc['columns'] if c['type'] == 'dim']
    olcu_kolonlari = [c for c in sonuc['columns'] if c['type'] != 'dim']
    if not olcu_kolonlari:
        return None
    kolonlar.append(olcu_kolonlari[0])
    if gb:
        if ana_tip != 'pct':
            kolonlar.append({'key': '_pay', 'label': 'Pay', 'type': 'pct'})
        kolonlar.append({'key': '_onceki', 'label': 'Bir önceki dönem' if zaman else 'Önceki eşit dönem', 'type': ana_tip})
        kolonlar.append({'key': '_degisim', 'label': 'Fark (puan)' if ana_tip == 'pct' else 'Değişim', 'type': 'puan' if ana_tip == 'pct' else 'degisim'})
    kolonlar += olcu_kolonlari[1:]
    toplam = dict(sonuc['totals'])
    if gb and ana_tip != 'pct':
        toplam['_pay'] = 100.0
    if not zaman and isinstance(onceki_toplam, (int, float)):
        toplam['_onceki'] = onceki_toplam

    t = sonuc['totals'].get(ana)
    degisim = None
    if isinstance(t, (int, float)) and isinstance(onceki_toplam, (int, float)) and onceki_toplam:
        degisim = (t - onceki_toplam) if ana_tip == 'pct' else 100.0 * (t - onceki_toplam) / abs(onceki_toplam)
    en_yuksek = max(satirlar, key=lambda r: r.get(ana) if isinstance(r.get(ana), (int, float)) else float('-inf'), default=None)
    return {'kolonlar': kolonlar, 'satirlar': satirlar, 'toplam': toplam, 'zaman': zaman, 'ana': ana, 'ana_tip': ana_tip,
            'ozet': {'toplam': t, 'onceki': onceki_toplam, 'degisim': degisim, 'kalem': len(satirlar),
                     'en_yuksek': (en_yuksek.get('d0'), en_yuksek.get(ana)) if en_yuksek and gb else None,
                     'olcu_adi': olcu_kolonlari[0]['label']}}


def _ozet_kutulari(ws, r, ozet, tip):
    """Baslik altinda yan yana ozet degerleri (etiket ustte, deger altta)."""
    kutular = [(ozet['olcu_adi'], ozet['toplam'], tip),
               ('Önceki eşit dönem', ozet['onceki'], tip),
               ('Değişim' if tip != 'pct' else 'Fark (puan)', ozet['degisim'], 'degisim' if tip != 'pct' else 'puan')]
    if ozet.get('en_yuksek'):
        kutular.append((f"En yüksek: {ozet['en_yuksek'][0]}", ozet['en_yuksek'][1], tip))
    for j, (ad, v, t) in enumerate(kutular):
        col = 1 + j * 2
        a = ws.cell(row=r, column=col, value=ad)
        a.font = F_ALT
        b = ws.cell(row=r + 1, column=col, value=v if isinstance(v, (int, float)) else '—')
        b.alignment = Alignment(horizontal='left')
        for satir in (r, r + 1):                                   # her kutu iki sutun genisliginde
            ws.merge_cells(start_row=satir, start_column=col, end_row=satir, end_column=col + 1)
        b.font = Font(name=YAZI, size=14, bold=True, color=(('15803D' if v > 0 else 'B91C1C') if t in ('degisim', 'puan') and isinstance(v, (int, float)) else MUREKKEP))
        if isinstance(v, (int, float)):
            b.number_format = _bicim(t, v)
    ws.row_dimensions[r + 1].height = 22
    return r + 3


def sorgu_excel(istek, baslik='Analiz'):
    """Kart / odak / analiz sayfasi disa aktarimi: ozet + karsilastirmali tablo + grafik."""
    if istek.get('pivot_by'):                                    # pivot: eski duz tablo
        sonuc = analytics.sorgu({**istek, 'limit': 0})
        wb = Workbook()
        _bilgi_sayfasi(wb, istek, baslik)
        _sayfa_yaz(wb.create_sheet('Veri'), sonuc['columns'], sonuc['rows'], None, baslik=baslik, alt=(_aralik_yazisi(istek),))
        return _kaydet(wb)
    wb = Workbook()
    ws = wb.active
    ws.title = _sayfa_adi(wb, baslik)
    _kart_sayfasi(ws, istek, baslik)
    return _kaydet(wb)


def _kart_sayfasi(ws, istek, baslik, not_=None, yorum=()):
    a = analiz_tablosu(istek)
    alt = [f"Dönem: {_aralik_yazisi(istek)}" + (f"  ·  karşılaştırma: {_aralik_yazisi(dict(zip(('date_from', 'date_to'), analytics.onceki_donem(istek['date_from'], istek['date_to']))))}"
                                                 if istek.get('date_from') and istek.get('date_to') else ''),
           _filtre_yazisi(istek), f"Üretildi: {datetime.now().strftime('%d.%m.%Y %H:%M')}"]
    r = _ust_bant(ws, baslik[:1].upper() + baslik[1:], alt, max(len(a['kolonlar']) if a else 4, 7))
    if not a:
        ws.cell(row=r, column=1, value='Veri yok.').font = F_HUCRE
        return None
    r = _ozet_kutulari(ws, r, a['ozet'], a['ana_tip'])
    for metin, etiket in ([(c, 'Yorum') for c in yorum] + ([(not_, 'Not')] if not_ else [])):
        ws.cell(row=r, column=1, value=etiket).font = F_ETIKET
        c = ws.cell(row=r, column=2, value=metin)
        c.font, c.alignment = F_HUCRE, Alignment(wrap_text=False, vertical='top')
        r += 1
    if yorum or not_:
        r += 1
    if istek.get('group_by'):
        _sayfa_yaz(ws, a['kolonlar'], a['satirlar'], a['toplam'], r0=r,
                   grafik={'olcu_key': a['ana'], 'zaman': a['zaman']})
    return a


def satirlar_excel(istek, baslik='Hareket listesi'):
    # Excel tarafinda tum satirlar: 5000'lik sayfalarla
    tum, offset = [], 0
    kolonlar = None
    while True:
        s = analytics.satirlar({**istek, 'limit': 5000, 'offset': offset})
        kolonlar = s['columns']
        tum += s['rows']
        offset += 5000
        if offset >= s['total'] or offset >= 200000:
            break
    secili = istek.get('columns')
    if secili:
        kolonlar = [k for k in kolonlar if k['key'] in secili]
    for r in tum:
        for k in kolonlar:
            if k['type'] == 'date' and r.get(k['key']):
                r[k['key']] = _tarih(r[k['key']])
    wb = Workbook()
    _bilgi_sayfasi(wb, istek, baslik)
    toplam = {k['key']: sum((r.get(k['key']) or 0) for r in tum) for k in kolonlar
              if k['type'] in ('tl', 'usd') or k['key'] == 'qty'}
    _sayfa_yaz(wb.create_sheet('Satırlar'), kolonlar, tum, toplam, baslik=baslik,
               alt=(f"Dönem: {_aralik_yazisi(istek)}", _filtre_yazisi(istek), f'{len(tum):,} satır'.replace(',', '.')))
    return _kaydet(wb)


def stok_excel(istek):
    s = analytics.stok({**istek, 'limit': 0})
    kolonlar = [
        {'key': 'item', 'label': 'Ürün', 'type': 'text'}, {'key': 'item_code', 'label': 'Kod', 'type': 'text'},
        {'key': 'brand', 'label': 'Marka', 'type': 'text'}, {'key': 'category', 'label': 'Kategori', 'type': 'text'},
        {'key': 'qty', 'label': 'Eldeki adet', 'type': 'num'},
        {'key': 'unit_cost_tl', 'label': 'Son alış birim (TL)', 'type': 'tl'},
        {'key': 'value_tl', 'label': 'Değer (TL)', 'type': 'tl'}, {'key': 'value_usd', 'label': 'Değer (USD)', 'type': 'usd'},
        {'key': 'last_purchase', 'label': 'Son alış', 'type': 'date'}, {'key': 'last_sale', 'label': 'Son satış', 'type': 'date'},
        {'key': 'days_since_move', 'label': 'Hareketsiz (gün)', 'type': 'num'},
        {'key': 'sold_recent', 'label': f"Son {s['velocity_days']} gün satış", 'type': 'num'},
        {'key': 'cover_days', 'label': 'Kaç gün yeter', 'type': 'num'},
    ]
    satirlar = [{**r, 'last_purchase': _tarih(r.get('last_purchase')), 'last_sale': _tarih(r.get('last_sale'))} for r in s['rows']]
    wb = Workbook()
    _bilgi_sayfasi(wb, {**istek, 'date_to': s['as_of'], 'date_from': None}, f"Stok durumu — {_tarih(s['as_of'])}")
    _sayfa_yaz(wb.create_sheet('Stok'), kolonlar, satirlar,
               {'qty': s['totals']['qty'], 'value_tl': s['totals']['value_tl'], 'value_usd': s['totals']['value_usd']},
               baslik=f"Stok durumu — {_tarih(s['as_of'])}", alt=(_filtre_yazisi(istek),))
    return _kaydet(wb)


def donem_excel(istek):
    d = analytics.donem_raporu(istek)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Özet'
    _sayfa_hazirla(ws, yatay=False)
    r = _ust_bant(ws, 'Dönem raporu', (f"Dönem: {_aralik_yazisi(d)}", _filtre_yazisi(istek),
                                       f"Üretildi: {datetime.now().strftime('%d.%m.%Y %H:%M')}"), 2)
    t = d['totals']
    ozet = [
        ('Dönem başı stok değeri', t['open_value_tl'], 'tl'), ('Alışlar', t['buy_tl'], 'tl'),
        ('Alışlar (USD)', t['buy_usd'], 'usd'), ('Alış faturası', t['buy_invoices'], 'num'),
        ('Satışlar', t['sell_tl'], 'tl'), ('Satışlar (USD)', t['sell_usd'], 'usd'),
        ('Satış faturası', t['sell_invoices'], 'num'), ('Müşteri sayısı', t['clients'], 'num'),
        ('Brüt kâr (FIFO)', t['profit_tl'], 'tl'), ('Kâr marjı', t['margin'], 'pct'),
        ('Dönem sonu stok değeri', t['close_value_tl'], 'tl'), ('Dönem sonu stok değeri (USD)', t['close_value_usd'], 'usd'),
    ]
    for i, (k, v, tip) in enumerate(ozet):
        a, b = ws.cell(row=r + i, column=1, value=k), ws.cell(row=r + i, column=2, value=v)
        a.font, b.font = F_ETIKET, F_HUCRE
        if isinstance(v, (int, float)):
            b.number_format = _bicim(tip, v)
        for c in (a, b):
            c.border = Border(bottom=INCE)
            if i % 2:
                c.fill = D_SERIT
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 22

    kolonlar = [
        {'key': 'group', 'label': d['group_label'], 'type': 'text'},
        {'key': 'open_qty', 'label': 'Dönem başı adet', 'type': 'num'},
        {'key': 'open_value_tl', 'label': 'Dönem başı değer', 'type': 'tl'},
        {'key': 'buy_qty', 'label': 'Alınan adet', 'type': 'num'}, {'key': 'buy_tl', 'label': 'Alış', 'type': 'tl'},
        {'key': 'buy_usd', 'label': 'Alış (USD)', 'type': 'usd'},
        {'key': 'sell_qty', 'label': 'Satılan adet', 'type': 'num'}, {'key': 'sell_tl', 'label': 'Satış', 'type': 'tl'},
        {'key': 'sell_usd', 'label': 'Satış (USD)', 'type': 'usd'},
        {'key': 'profit_tl', 'label': 'Brüt kâr', 'type': 'tl'}, {'key': 'margin', 'label': 'Marj', 'type': 'pct'},
        {'key': 'close_qty', 'label': 'Dönem sonu adet', 'type': 'num'},
        {'key': 'close_value_tl', 'label': 'Dönem sonu değer', 'type': 'tl'},
    ]
    _sayfa_yaz(wb.create_sheet(_sayfa_adi(wb, d['group_label'])), kolonlar, d['rows'],
               {k['key']: sum((r.get(k['key']) or 0) for r in d['rows']) for k in kolonlar if k['type'] in ('tl', 'usd', 'num')},
               baslik=f"Dönem raporu — {d['group_label']}", alt=(f"Dönem: {_aralik_yazisi(d)}",))
    for ad, anahtar in (('En çok satılan müşteriler', 'top_customers'), ('En çok alınan tedarikçiler', 'top_suppliers'),
                        ('En çok satılan ürünler', 'top_items_sold'), ('En çok alınan ürünler', 'top_items_bought')):
        _sayfa_yaz(wb.create_sheet(_sayfa_adi(wb, ad)), [{'key': 'd0', 'label': 'Ad', 'type': 'text'},
                                                         {'key': 'm_net_tl', 'label': 'Tutar', 'type': 'tl'},
                                                         {'key': 'm_net_usd', 'label': 'Tutar (USD)', 'type': 'usd'}],
                   d[anahtar], baslik=ad, alt=(f"Dönem: {_aralik_yazisi(d)}",))
    return _kaydet(wb)


def pano_excel(govde):
    """Analiz Studyosu: ilk sayfa ozet (her kartin toplami, degisimi, yorumu, sayfasina baglanti),
    sonra her kart kendi sayfasinda (ozet + karsilastirmali tablo + grafik).
    govde: {'ad', 'not', 'kartlar': [{'baslik', 'request', 'not', 'yorum': [..]}]}"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Özet'
    _sayfa_hazirla(ws)
    kartlar = govde.get('kartlar') or []
    r = _ust_bant(ws, govde.get('ad') or 'Analiz',
                  (f"{len(kartlar)} analiz kartı", f"Üretildi: {datetime.now().strftime('%d.%m.%Y %H:%M')}"), 6)
    if govde.get('not'):
        ws.cell(row=r, column=1, value='Analiz notu').font = F_ETIKET
        c = ws.cell(row=r, column=2, value=govde['not'])
        c.font, c.alignment = F_HUCRE, Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = max(18, 15 * (1 + len(govde['not']) // 110))
        r += 2

    basliklar = ['Analiz', 'Dönem', 'Toplam', 'Önceki eşit dönem', 'Değişim', 'Sayfa']
    for j, b in enumerate(basliklar, 1):
        c = ws.cell(row=r, column=j, value=b)
        c.fill, c.font = D_KOLON, F_KOLON
        c.alignment = Alignment(horizontal='left' if j in (1, 2, 6) else 'right', vertical='center')
    ws.row_dimensions[r].height = 24
    ozet_satiri = r + 1
    r = ozet_satiri + len(kartlar) + 2

    yorum_bas = r
    for i, k in enumerate(kartlar, 1):
        istek = k.get('request') or {}
        baslik = k.get('baslik') or 'Kart'
        ad = _sayfa_adi(wb, f'{i}. {baslik}')
        a = _kart_sayfasi(wb.create_sheet(ad), istek, baslik, k.get('not'), k.get('yorum') or [])
        o = (a or {}).get('ozet') or {}
        tip = (a or {}).get('ana_tip', 'num')
        satir = ozet_satiri + i - 1
        degerler = [baslik[:1].upper() + baslik[1:], _aralik_yazisi(istek), o.get('toplam'), o.get('onceki'), o.get('degisim')]
        for j, v in enumerate(degerler, 1):
            c = ws.cell(row=satir, column=j, value=v)
            c.font, c.border = F_HUCRE, Border(bottom=INCE)
            if i % 2 == 0:
                c.fill = D_SERIT
            if isinstance(v, (int, float)):
                t = ('degisim' if tip != 'pct' else 'puan') if j == 5 else tip
                c.number_format = _bicim(t, v)
                if j == 5:
                    c.font = F_POZ if v > 0 else F_NEG if v < 0 else F_HUCRE
        git = ws.cell(row=satir, column=6, value='Aç →')
        git.hyperlink = f"#'{ad}'!A1"
        git.font = Font(name=YAZI, size=10, color=TURKUAZ_KOYU, underline='single')
        git.border = Border(bottom=INCE)

        # yorumlar ozet sayfasinin altinda
        ws.cell(row=r, column=1, value=f'{i}. {baslik[:1].upper() + baslik[1:]}').font = F_ETIKET
        r += 1
        for cumle in k.get('yorum') or []:
            c = ws.cell(row=r, column=1, value=f'•  {cumle}')
            c.font, c.alignment = F_HUCRE, Alignment(wrap_text=True, vertical='top')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            ws.row_dimensions[r].height = max(15, 15 * (1 + len(cumle) // 120))
            r += 1
        if k.get('not'):
            c = ws.cell(row=r, column=1, value=f'Not: {k["not"]}')
            c.font, c.alignment = Font(name=YAZI, size=10, italic=True, color=GRI), Alignment(wrap_text=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            r += 1
        r += 1
    if kartlar:
        ws.cell(row=yorum_bas - 1, column=1, value='Ne söylüyor?').font = Font(name=YAZI, size=12, bold=True, color=MUREKKEP)
    for col, w in zip('ABCDEF', (46, 26, 18, 18, 12, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = None
    return _kaydet(wb)
